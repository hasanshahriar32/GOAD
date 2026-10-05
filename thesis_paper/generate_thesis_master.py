#!/usr/bin/env python3
"""
Generate the complete, publication-grade thesis Word Document (thesis.docx)
using native python-docx frontmatter construction combined with Pandoc body
compilation and docxcompose stitching.
"""

import os
import sys
import subprocess
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docxcompose.composer import Composer

def set_cell_margins(cell, top=20, bottom=20, left=60, right=60):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def remove_table_borders(table):
    tblPr = table._tbl.tblPr
    for child in list(tblPr):
        if child.tag.endswith('tblBorders'):
            tblPr.remove(child)
    tblBorders = OxmlElement('w:tblBorders')
    for b in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        node = OxmlElement(f'w:{b}')
        node.set(qn('w:val'), 'none')
        tblBorders.append(node)
    tblPr.append(tblBorders)

def add_toc_field(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement('w:fldChar')
    fld1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld2 = OxmlElement('w:fldChar')
    fld2.set(qn('w:fldCharType'), 'separate')
    fld3 = OxmlElement('w:fldChar')
    fld3.set(qn('w:fldCharType'), 'end')
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    run._r.append(fld3)

def build_frontmatter(output_path):
    doc = docx.Document()
    
    # Configure A4 margins: Left 1.25", others 1.0"
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(1.0)
    sec.bottom_margin = Inches(1.0)
    sec.left_margin = Inches(1.25)
    sec.right_margin = Inches(1.0)

    def p(text="", size=11, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=4, line_spacing=1.15):
        par = doc.add_paragraph()
        par.alignment = align
        par.paragraph_format.space_before = Pt(space_before)
        par.paragraph_format.space_after = Pt(space_after)
        par.paragraph_format.line_spacing = line_spacing
        if text:
            r = par.add_run(text)
            r.font.name = "Times New Roman"
            r.font.size = Pt(size)
            r.bold = bold
            r.italic = italic
        return par

    # ─────────────────────────────────────────────────────────────
    # PAGE 1: COVER PAGE
    # ─────────────────────────────────────────────────────────────
    p("CertGraph: Heterogeneous Graph Attention Networks for Active Directory Certificate Services Vulnerability Detection and Autonomous Defense",
      size=17, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
      
    p("Course Code: ECE 452        Course Title: Project and Thesis",
      size=11.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
      
    p("Submitted By—", size=11.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    
    # Student IDs with Level 4, Semester II
    p("Student ID: 2002126        Level: 4, Semester: II", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    p("Student ID: 2002138        Level: 4, Semester: II", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    p("Student ID: 2102151        Level: 4, Semester: II", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    
    # HSTU Logo
    logo_par = doc.add_paragraph()
    logo_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_par.paragraph_format.space_before = Pt(4)
    logo_par.paragraph_format.space_after = Pt(10)
    if os.path.exists("figures/hstu_logo.png"):
        logo_par.add_run().add_picture("figures/hstu_logo.png", width=Inches(1.1))
        
    p("Submitted To—", size=11.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    p("Department of Electronics and Communication Engineering", size=12.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    p("in partial fulfillment of the requirements for the degree of", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    p("Bachelor of Science in Electronics and Communication Engineering", size=11.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    p("Hajee Mohammad Danesh Science and Technology University (HSTU)", size=12.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    p("Dinajpur-5200, Bangladesh", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    p("October, 2026", size=11.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    
    doc.add_page_break()

    # ─────────────────────────────────────────────────────────────
    # PAGE 2: CERTIFICATE
    # ─────────────────────────────────────────────────────────────
    p("Certificate", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)
    
    p("This is to certify that the thesis work entitled “CertGraph: Heterogeneous Graph Attention Networks for Active Directory Certificate Services Vulnerability Detection and Autonomous Defense” is carried by the following ID numbers: 2002126, 2002138, 2102151. To the fullest extent of our knowledge, we assert that this undertaking is an authentic and original contribution to the field. We certify that this thesis has not been previously submitted for the award of any other degree or diploma at this or any other institution.",
      size=11.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=16, line_spacing=1.2)
      
    p("Signed by the Final Examining committee:", size=11.5, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=16)
    
    # 2-column signature table
    tbl_cert = doc.add_table(rows=7, cols=2)
    remove_table_borders(tbl_cert)
    tbl_cert.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_cert.autofit = False
    
    cert_data = [
        ("........................................", "........................................"),
        ("Chairman", "Supervisor"),
        ("Examination Committee", "Department of ECE, HSTU"),
        ("........................................", "........................................"),
        ("External Member", "Co-Supervisor"),
        ("Examination Committee", "Department of ECE, HSTU"),
        ("........................................\nInternal Member\nExamination Committee", "")
    ]
    
    for r_idx, (c1, c2) in enumerate(cert_data):
        row = tbl_cert.rows[r_idx]
        row.cells[0].width = Inches(2.9)
        row.cells[1].width = Inches(2.9)
        set_cell_margins(row.cells[0])
        set_cell_margins(row.cells[1])
        
        p1 = row.cells[0].paragraphs[0]
        p1.paragraph_format.line_spacing = 1.05
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(c1)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(10.5)
        if "Chairman" in c1 or "External" in c1 or "Internal" in c1:
            r1.bold = True
            
        p2 = row.cells[1].paragraphs[0]
        p2.paragraph_format.line_spacing = 1.05
        p2.paragraph_format.space_after = Pt(2)
        r2 = p2.add_run(c2)
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(10.5)
        if "Supervisor" in c2 or "Co-Supervisor" in c2:
            r2.bold = True

    doc.add_page_break()

    # ─────────────────────────────────────────────────────────────
    # PAGE 3: CANDIDATE'S DECLARATION
    # ─────────────────────────────────────────────────────────────
    p("Candidate's Declaration", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    
    p("We hereby declare that the research work presented in this thesis entitled “CertGraph: Heterogeneous Graph Attention Networks for Active Directory Certificate Services Vulnerability Detection and Autonomous Defense” is the outcome of an original investigation conducted by us under the supervision of the Department of Electronics and Communication Engineering, Hajee Mohammad Danesh Science and Technology University (HSTU), Dinajpur-5200, Bangladesh.",
      size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8, line_spacing=1.15)
      
    p("We further solemnly declare that:", size=11, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4)
    
    p("1.  This work, or any part thereof, has not been submitted previously to any university or institution for the award of any degree, diploma, or other academic qualification.",
      size=10.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=2)
    p("2.  All material, concepts, and algorithms taken from the published or unpublished work of others have been fully and properly acknowledged and cited in accordance with standard academic referencing protocols.",
      size=10.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=2)
    p("3.  All synthetic datasets, forensic auditing scripts, empirical benchmark routines, and graph learning models described herein were constructed with rigorous adherence to ethical scientific standards and academic integrity.",
      size=10.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=12)
      
    tbl_decl = doc.add_table(rows=4, cols=2)
    remove_table_borders(tbl_decl)
    tbl_decl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_decl.autofit = False
    
    decl_data = [
        ("Date: October, 2026", "Student ID: 2002126"),
        ("Place: HSTU, Dinajpur", "Student ID: 2002138"),
        ("", "Student ID: 2102151"),
        ("", "Department of ECE, HSTU")
    ]
    
    for r_idx, (c1, c2) in enumerate(decl_data):
        row = tbl_decl.rows[r_idx]
        row.cells[0].width = Inches(2.9)
        row.cells[1].width = Inches(2.9)
        set_cell_margins(row.cells[0])
        set_cell_margins(row.cells[1])
        
        p1 = row.cells[0].paragraphs[0]
        p1.paragraph_format.line_spacing = 1.05
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(c1)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(10.5)
        if "Date" in c1 or "Place" in c1:
            r1.bold = True
            
        p2 = row.cells[1].paragraphs[0]
        p2.paragraph_format.line_spacing = 1.05
        p2.paragraph_format.space_after = Pt(2)
        r2 = p2.add_run(c2)
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(10.5)

    doc.add_page_break()

    # ─────────────────────────────────────────────────────────────
    # PAGE 4: DEDICATION
    # ─────────────────────────────────────────────────────────────
    p("", space_after=140)
    p("Dedication", size=17, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=20)
    p("This thesis is dedicated to our beloved parents,\nwhose endless sacrifices, prayers, and unconditional love\nhave been the guiding light of our lives.",
      size=12, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16, line_spacing=1.3)
    p("And to all our teachers and mentors,\nwho inspired our passion for computer science, cybersecurity, and scientific discovery.",
      size=12, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0, line_spacing=1.3)
      
    doc.add_page_break()

    # ─────────────────────────────────────────────────────────────
    # PAGE 5: ACKNOWLEDGEMENTS
    # ─────────────────────────────────────────────────────────────
    p("Acknowledgements", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)
    
    p("First and foremost, all praises are due to Almighty Allah, the Most Merciful and Most Beneficent, who bestowed upon us the health, strength, patience, and intellect required to complete this project and thesis research successfully.",
      size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)
    p("We express our profound gratitude, respect, and deepest indebtedness to our respected thesis supervisor and co-supervisor in the Department of Electronics and Communication Engineering, Hajee Mohammad Danesh Science and Technology University (HSTU), Dinajpur. Their exemplary guidance, insightful suggestions, constant encouragement, and critical academic reviews were invaluable throughout the formulation, mathematical derivation, experimental validation, and manuscript preparation of this research.",
      size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)
    p("We are deeply thankful to the Chairman and all distinguished faculty members of the Department of Electronics and Communication Engineering, HSTU, for providing a vibrant academic environment, high-quality computational facilities, and continuous moral support during our undergraduate curriculum.",
      size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)
    p("Our sincere gratitude goes to the global open-source cybersecurity and machine learning research communities. In particular, we acknowledge the pioneering work of Will Schroeder and Lee Christensen (SpecterOps) for uncovering the ADCS attack surface; the creators of BloodHound, SharpHound, and Certipy for their foundational offensive graph tools; the developers of the Game of Active Directory (GOAD) laboratory; and the core contributors of PyTorch Geometric and NetworkX.",
      size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)
    p("Finally, we owe an immeasurable debt of gratitude to our parents and families for their unending sacrifices, patience, and blessings throughout our university education. We also express warm appreciation to our batchmates, lab peers, and friends whose intellectual discussions, camaraderie, and encouragement enriched every stage of this thesis journey.",
      size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=14)
      
    tbl_ack = doc.add_table(rows=2, cols=2)
    remove_table_borders(tbl_ack)
    tbl_ack.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_ack.autofit = False
    
    ack_row0 = tbl_ack.rows[0]
    ack_row0.cells[0].width = Inches(2.9)
    ack_row0.cells[1].width = Inches(2.9)
    p_a1 = ack_row0.cells[0].paragraphs[0]
    p_a1.add_run("HSTU, Dinajpur\nOctober, 2026").bold = True
    p_a2 = ack_row0.cells[1].paragraphs[0]
    p_a2.add_run("Student ID: 2002126\nStudent ID: 2002138\nStudent ID: 2102151")
    
    ack_row1 = tbl_ack.rows[1]
    ack_row1.cells[0].width = Inches(2.9)
    ack_row1.cells[1].width = Inches(2.9)
    p_a4 = ack_row1.cells[1].paragraphs[0]
    p_a4.add_run("Department of ECE, HSTU")
    
    doc.add_page_break()

    # ─────────────────────────────────────────────────────────────
    # PAGE 6: ABSTRACT
    # ─────────────────────────────────────────────────────────────
    p("Abstract", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    
    p("Active Directory Certificate Services (ADCS) forms the cryptographic backbone of enterprise identity, issuing credentials for domain authentication and single sign-on. Misconfigured certificate templates, access control lists (ACLs), and issuance policies introduce systemic privilege escalation vectors (ESC1--ESC15+) that enable unprivileged accounts to compromise entire Active Directory domains. Signature-based scanners (e.g., Certipy, BloodHound) evaluate flags in isolation and incur high traversal overhead on large topologies. In this thesis, we present CertGraph, a heterogeneous Graph Attention Network (Hetero-GAT) that models enterprise identity as a typed multigraph G = (V, E, T_V, T_E) to detect structural escalation paths across core vectors (ESC1, ESC2, ESC3, ESC4, ESC9, and ESC13).",
      size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, line_spacing=1.15)
      
    p("Grounding our evaluation in the security machine learning framework of Arp et al. (USENIX Security 2022), we identify and resolve three systemic pitfalls in synthetic identity datasets: positional index leakage, baseline feature asymmetry, and rule saturation. We prove mathematically (Theorem 1) that residual skip-connections are indispensable for preventing representation collapse on identity graphs; omitting them makes template representations independent of input flags, reducing classification Macro-F1 from 0.9986 to 0.4768. On an audited 700-environment enterprise benchmark evaluated under 5-fold cross-validation with Nadeau-Bengio corrected variance, CertGraph achieves a Macro-F1 of 0.9986, substantially outperforming flat classifiers (0.8600) and signature heuristics (0.7791).",
      size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, line_spacing=1.15)
      
    p("However, when evaluated against an adversarial Zero-Shot Hard Negative suite (n = 63) containing benign templates with dangerous flags, pure neural models collapse to 1.59% accuracy due to shortcut learning on isolated attributes, whereas symbolic graph traversal achieves 84.13%. To resolve this dichotomy, we formulate a Two-Tier Neuro-Symbolic architecture pairing rapid GNN candidate screening (filtering >89% of benign templates) with a localized symbolic verification oracle to eliminate false positives by construction. Finally, we model identity privilege revocation as a Stackelberg security game, prove its NP-hardness via reduction from Directed Multiway Cut, and demonstrate a greedy capacity-disruption heuristic that severs up to 68.6% of compromise paths under bounded operational budgets (p < 10^-8).",
      size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=10, line_spacing=1.15)
      
    par_kw = p("", size=11, space_after=0)
    r_kwh = par_kw.add_run("Keywords: ")
    r_kwh.bold = True
    r_kwh.font.name = "Times New Roman"
    r_kwt = par_kw.add_run("Active Directory Certificate Services (ADCS), Graph Attention Networks, Heterogeneous Graphs, Identity and Access Management, Neuro-Symbolic Security, Autonomous Cyber Defense.")
    r_kwt.font.name = "Times New Roman"
    
    doc.add_page_break()

    # ─────────────────────────────────────────────────────────────
    # PAGE 7: TABLE OF CONTENTS
    # ─────────────────────────────────────────────────────────────
    p("Table of Contents", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16)
    toc_p = doc.add_paragraph()
    add_toc_field(toc_p)
    
    doc.add_page_break()

    doc.save(output_path)
    print(f"[✓] Created clean frontmatter: '{output_path}'")

def main():
    dir_path = os.path.dirname(os.path.abspath(__file__))
    os.chdir(dir_path)

    front_docx = "/tmp/frontmatter.docx"
    body_docx = "/tmp/body.docx"
    final_docx = "thesis.docx"

    print("[1/4] Generating pixel-perfect frontmatter via python-docx...")
    build_frontmatter(front_docx)

    print("[2/4] Compiling thesis body chapters, theorems, and bibliography via Pandoc...")
    cmd = [
        "pandoc",
        "chapters/ch01_introduction.tex",
        "chapters/ch02_background_threat.tex",
        "chapters/ch03_formal_methodology.tex",
        "chapters/ch04_forensic_audit.tex",
        "chapters/ch05_empirical_benchmarks.tex",
        "chapters/ch06_real_world_case_study.tex",
        "chapters/ch07_robustness_scalability.tex",
        "chapters/ch08_neuro_symbolic_hybrid.tex",
        "chapters/ch09_game_theoretic_defense.tex",
        "chapters/ch10_conclusion.tex",
        "proofs/theorem1_representation_collapse.tex",
        "proofs/theorem2_edge_blocking_nphardness.tex",
        "--citeproc",
        "--bibliography=references.bib",
        "--reference-doc=old-reference-paper-to-follow-design/final(corrected).docx",
        "-o", body_docx
    ]
    res = subprocess.run(cmd, capture_output=True, text=True)
    if res.returncode != 0:
        print(f"[-] Pandoc failed:\n{res.stderr}")
        sys.exit(1)

    print("[3/4] Refinement on body document (page breaks before chapters, centering figures)...")
    b_doc = docx.Document(body_docx)
    for p in b_doc.paragraphs:
        if p.style.name.startswith("Heading 1"):
            p.paragraph_format.page_break_before = True
        for r in p.runs:
            if 'graphic' in r._element.xml:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p.style.name.lower().startswith('caption'):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    b_doc.save(body_docx)

    print("[4/4] Stitching frontmatter and body via docxcompose...")
    master = docx.Document(front_docx)
    composer = Composer(master)
    composer.append(docx.Document(body_docx))
    composer.save(final_docx)

    desktop_path = "/home/hs32/Desktop/thesis.docx"
    shutil.copyfile(final_docx, desktop_path)

    size_mb = os.path.getsize(final_docx) / (1024 * 1024)
    print(f"[✓] Complete! Successfully generated '{final_docx}' ({size_mb:.2f} MB)")
    print(f"    Copied to '{desktop_path}'")

if __name__ == "__main__":
    main()
