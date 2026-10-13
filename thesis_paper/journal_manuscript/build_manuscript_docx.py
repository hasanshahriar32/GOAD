#!/usr/bin/env python3
"""
=============================================================================
CertGraph IEEE Journal Manuscript DOCX Builder (Native Office MathML OMML)
=============================================================================
Transforms the IEEE Transactions LaTeX manuscript (manuscript.tex) into a 
publication-grade Word document (.docx) with:
  - 100% Native Office MathML (OMML) equations
  - Formal journal title block, author affiliations, and abstract box
  - Classic academic booktabs tables (1.5pt top/bottom rules, no vertical lines)
  - Centered high-resolution vector figures scaled to text width
  - Formatted IEEE bibliography with hanging indents
  - Auto-copied to /home/hs32/Desktop/manuscript.docx
=============================================================================
"""

import os
import re
import shutil
import subprocess
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
SCRATCH_DIR = "/tmp/manuscript_docx_build"

def preprocess_latex_for_pandoc(src_path, dst_path):
    with open(src_path, "r", encoding="utf-8") as f:
        content = f.read()

    # 1. Expand IEEEPARstart
    content = re.sub(r"\\IEEEPARstart\{([^}]+)\}\{([^}]+)\}", r"\1\2", content)

    # 2. Convert table* and figure* to standard table and figure
    content = re.sub(r"\\begin\{table\*\}", r"\\begin{table}", content)
    content = re.sub(r"\\end\{table\*\}", r"\\end{table}", content)
    content = re.sub(r"\\begin\{figure\*\}", r"\\begin{figure}", content)
    content = re.sub(r"\\end\{figure\*\}", r"\\end{figure}", content)

    # 3. Handle algorithmic environment for pandoc
    # Convert algorithm to a styled box
    def alg_repl(m):
        caption = m.group(1)
        code = m.group(2)
        # Format code lines
        lines = code.strip().split("\n")
        formatted_lines = []
        for line in lines:
            line_str = line.strip()
            line_str = re.sub(r"\\STATE\s*", "    ", line_str)
            line_str = re.sub(r"\\REQUIRE\s*", "\\textbf{Input:} ", line_str)
            line_str = re.sub(r"\\ENSURE\s*", "\\textbf{Output:} ", line_str)
            line_str = re.sub(r"\\WHILE\{([^}]+)\}", r"\\textbf{while} $\1$ \\textbf{do}", line_str)
            line_str = re.sub(r"\\ENDWHILE", r"\\textbf{end while}", line_str)
            line_str = re.sub(r"\\IF\{([^}]+)\}", r"\\textbf{if} $\1$ \\textbf{then}", line_str)
            line_str = re.sub(r"\\ENDIF", r"\\textbf{end if}", line_str)
            line_str = re.sub(r"\\RETURN\s*", "\\textbf{return} ", line_str)
            formatted_lines.append(line_str)
        return (
            f"\\begin{{table}}[ht]\n"
            f"\\caption{{{caption}}}\n"
            f"\\begin{{tabularx}}{{\\columnwidth}}{{X}}\n"
            f"\\toprule\n"
            + " \\\\\n".join(formatted_lines) + " \\\\\n"
            f"\\bottomrule\n"
            f"\\end{{tabularx}}\n"
            f"\\end{{table}}"
        )

    content = re.sub(
        r"\\begin\{algorithm\}\[ht\]\s*\\caption\{([^}]+)\}\s*\\label\{[^}]+\}\s*\\begin\{algorithmic\}\[\d+\](.*?)(\\end\{algorithmic\}\s*\\end\{algorithm\})",
        alg_repl,
        content,
        flags=re.DOTALL
    )

    # 4. Standardize figure widths to textwidth
    content = re.sub(r"width=[\d\.]+\\columnwidth", r"width=1.0\\textwidth", content)
    content = re.sub(r"width=[\d\.]+\\textwidth", r"width=1.0\\textwidth", content)

    # 5. Remove balance command
    content = re.sub(r"\\balance", "", content)

    with open(dst_path, "w", encoding="utf-8") as f:
        f.write(content)

def set_cell_margins(cell, top=50, bottom=50, left=80, right=80):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_bottom_border(cell, sz="6", color="000000"):
    tcPr = cell._tc.get_or_add_tcPr()
    for child in list(tcPr):
        if child.tag.endswith('tcBorders'):
            tcPr.remove(child)
    tcBorders = OxmlElement('w:tcBorders')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), sz)
    bottom.set(qn('w:space'), '0')
    bottom.set(qn('w:color'), color)
    tcBorders.append(bottom)
    for b in ['top', 'left', 'right']:
        node = OxmlElement(f'w:{b}')
        node.set(qn('w:val'), 'none')
        tcBorders.append(node)
    tcPr.append(tcBorders)

def style_academic_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Check if table contains figures
    if any('w:drawing' in c._tc.xml for row in table.rows for c in row.cells):
        tblPr = table._tbl.tblPr
        for child in list(tblPr):
            if child.tag.endswith('tblBorders'):
                tblPr.remove(child)
        borders = parse_xml(r'<w:tblBorders %s><w:top w:val="none"/><w:bottom w:val="none"/><w:insideH w:val="none"/><w:insideV w:val="none"/><w:left w:val="none"/><w:right w:val="none"/></w:tblBorders>' % nsdecls('w'))
        tblPr.append(borders)
        return

    # Header Row
    hdr_row = table.rows[0]
    hdr_trPr = hdr_row._tr.get_or_add_trPr()
    hdr_trPr.append(parse_xml(r'<w:tblHeader %s/>' % nsdecls('w')))
    hdr_trPr.append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))
    
    for c_idx, cell in enumerate(hdr_row.cells):
        set_cell_margins(cell, top=60, bottom=60, left=80, right=80)
        set_cell_bottom_border(cell, sz="6", color="000000")
        for p in cell.paragraphs:
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.05
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.bold = True
                r.font.name = "Times New Roman"
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0, 0, 0)

    # Data Rows
    for r_idx in range(1, len(table.rows)):
        row = table.rows[r_idx]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(r'<w:cantSplit %s/>' % nsdecls('w')))
        for c_idx, cell in enumerate(row.cells):
            set_cell_margins(cell, top=40, bottom=40, left=80, right=80)
            for p in cell.paragraphs:
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                if c_idx > 0 and len(p.text.strip()) < 15:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                else:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(9.0)
                    r.font.color.rgb = RGBColor(0, 0, 0)

    # Booktabs Borders
    tblPr = table._tbl.tblPr
    for child in list(tblPr):
        if child.tag.endswith('tblBorders'):
            tblPr.remove(child)
    borders = parse_xml(
        r'<w:tblBorders %s>'
        r'  <w:top w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
        r'  <w:bottom w:val="single" w:sz="12" w:space="0" w:color="000000"/>'
        r'  <w:insideH w:val="none"/>'
        r'  <w:insideV w:val="none"/>'
        r'  <w:left w:val="none"/>'
        r'  <w:right w:val="none"/>'
        r'</w:tblBorders>' % nsdecls('w')
    )
    tblPr.append(borders)

def post_process_manuscript(raw_docx_path, output_docx_path):
    doc = docx.Document(raw_docx_path)

    # 1. Page Margins (Standard Letter, 1-inch all around)
    for sec in doc.sections:
        sec.page_width = Inches(8.5)
        sec.page_height = Inches(11.0)
        sec.top_margin = Inches(1.0)
        sec.bottom_margin = Inches(1.0)
        sec.left_margin = Inches(1.0)
        sec.right_margin = Inches(1.0)
        sec.header_distance = Inches(0.5)
        sec.footer_distance = Inches(0.5)

    # Set base font to Times New Roman
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(0, 0, 0)

    # 2. Iterate paragraphs
    in_abstract = False
    
    for p in doc.paragraphs:
        txt = p.text.strip()
        style_name = p.style.name if p.style else ""
        has_drawing = 'w:drawing' in p._p.xml
        has_display_math = 'm:oMathPara' in p._p.xml

        # Section Headings (Heading 1)
        if style_name.startswith("Heading 1"):
            in_abstract = False
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.first_line_indent = Pt(0)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)
                r.bold = True
                r.font.color.rgb = RGBColor(0, 0, 0)

        # Subsections (Heading 2)
        elif style_name.startswith("Heading 2"):
            in_abstract = False
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.first_line_indent = Pt(0)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(11)
                r.bold = True
                r.italic = True
                r.font.color.rgb = RGBColor(0, 0, 0)

        # Subsubsections (Heading 3)
        elif style_name.startswith("Heading 3"):
            in_abstract = False
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.first_line_indent = Pt(0)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(10.5)
                r.italic = True
                r.font.color.rgb = RGBColor(0, 0, 0)

        # Document Title
        elif style_name == "Title":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(12)
            p.paragraph_format.first_line_indent = Pt(0)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(18)
                r.bold = True
                r.font.color.rgb = RGBColor(0, 0, 0)

        # Authors
        elif style_name == "Author":
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(6)
            p.paragraph_format.first_line_indent = Pt(0)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(12)
                r.bold = True
                r.font.color.rgb = RGBColor(0, 0, 0)

        # Table Captions
        elif "table caption" in style_name.lower() or (txt.startswith("Table ") and len(txt) < 150):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(10)
                r.bold = True
                r.font.color.rgb = RGBColor(0, 0, 0)

        # Figure Captions
        elif "image caption" in style_name.lower() or (txt.startswith("Figure ") or txt.startswith("Fig. ")):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(12)
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0, 0, 0)

        # Figures (Images)
        elif has_drawing:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.keep_with_next = True

            # Scale to 6.20 inches width
            max_w_emu = int(6.20 * 914400)
            max_h_emu = int(6.80 * 914400)
            for ext in p._p.xpath('.//wp:extent'):
                cx = int(ext.get('cx', 0))
                cy = int(ext.get('cy', 0))
                if cx > int(2.0 * 914400):
                    ratio = cy / cx
                    new_cx = max_w_emu
                    new_cy = int(new_cx * ratio)
                    if new_cy > max_h_emu:
                        new_cy = max_h_emu
                        new_cx = int(new_cy / ratio)
                    ext.set('cx', str(new_cx))
                    ext.set('cy', str(new_cy))
            for a_ext in p._p.xpath('.//a:ext'):
                cx = int(a_ext.get('cx', 0))
                cy = int(a_ext.get('cy', 0))
                if cx > int(2.0 * 914400):
                    ratio = cy / cx
                    new_cx = max_w_emu
                    new_cy = int(new_cx * ratio)
                    if new_cy > max_h_emu:
                        new_cy = max_h_emu
                        new_cx = int(new_cy / ratio)
                    a_ext.set('cx', str(new_cx))
                    a_ext.set('cy', str(new_cy))

        # Math Formulas
        elif has_display_math:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = 1.15

        # Bibliography
        elif style_name == "Bibliography":
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.left_indent = Inches(0.35)
            p.paragraph_format.first_line_indent = -Inches(0.35)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.15
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(9.5)
                r.font.color.rgb = RGBColor(0, 0, 0)

        # Body Text / First Paragraph
        else:
            # Check for abstract
            if "abstract" in txt.lower() and len(txt) < 30:
                in_abstract = True
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(4)
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(11)
                    r.bold = True
                continue

            if in_abstract:
                p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                p.paragraph_format.left_indent = Inches(0.4)
                p.paragraph_format.right_indent = Inches(0.4)
                p.paragraph_format.first_line_indent = Pt(0)
                p.paragraph_format.line_spacing = 1.2
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.space_after = Pt(6)
                for r in p.runs:
                    r.font.name = "Times New Roman"
                    r.font.size = Pt(10)
                    r.font.color.rgb = RGBColor(0, 0, 0)
                continue

            # Standard body paragraphs
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Inches(0.25)
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.35
            for r in p.runs:
                r.font.name = "Times New Roman"
                r.font.size = Pt(11)
                r.font.color.rgb = RGBColor(0, 0, 0)

    # 3. Format Tables
    for tbl in doc.tables:
        style_academic_table(tbl)

    doc.save(output_docx_path)
    print(f"[✓] Successfully saved styled docx: '{output_docx_path}'")

def main():
    print("==================================================================")
    print("   COMPILING IEEE JOURNAL MANUSCRIPT DOCX (OFFICE MATH ML)        ")
    print("==================================================================")

    os.makedirs(SCRATCH_DIR, exist_ok=True)
    shutil.copytree(os.path.join(SCRIPT_DIR, "figures"), os.path.join(SCRATCH_DIR, "figures"), dirs_exist_ok=True)
    tex_src = os.path.join(SCRIPT_DIR, "manuscript.tex")
    preprocessed_tex = os.path.join(SCRATCH_DIR, "manuscript_clean.tex")
    raw_docx = os.path.join(SCRATCH_DIR, "manuscript_raw.docx")
    final_docx = os.path.join(SCRIPT_DIR, "manuscript.docx")
    desktop_docx = "/home/hs32/Desktop/manuscript.docx"

    print("[Step 1/3] Preprocessing manuscript.tex for pandoc...")
    preprocess_latex_for_pandoc(tex_src, preprocessed_tex)

    print("[Step 2/3] Compiling to Word with Pandoc (Pure OMML math)...")
    pandoc_cmd = [
        "pandoc",
        preprocessed_tex,
        f"--resource-path={SCRATCH_DIR}:{SCRIPT_DIR}:{os.path.join(SCRIPT_DIR, 'figures')}",
        "--number-sections",
        "--citeproc",
        f"--bibliography={os.path.join(SCRIPT_DIR, 'references.bib')}",
        f"--csl={os.path.join(SCRIPT_DIR, 'ieee.csl')}",
        "-o", raw_docx
    ]
    res = subprocess.run(pandoc_cmd, cwd=SCRATCH_DIR, capture_output=True, text=True)
    if res.returncode != 0:
        print(f"[-] Pandoc compilation error:\n{res.stderr}")
        sys.exit(1)
    print(f"[✓] Pandoc compiled successfully: '{raw_docx}'")

    print("[Step 3/3] Styling docx with academic booktabs and journal typography...")
    post_process_manuscript(raw_docx, final_docx)

    shutil.copyfile(final_docx, desktop_docx)
    size_mb = os.path.getsize(final_docx) / (1024 * 1024)
    print("==================================================================")
    print(f" [✓] COMPLETED SUCCESSFULLY: '{final_docx}' ({size_mb:.2f} MB)")
    print(f" [✓] COPIED TO DESKTOP: '{desktop_docx}'")
    print("==================================================================")

if __name__ == "__main__":
    main()
