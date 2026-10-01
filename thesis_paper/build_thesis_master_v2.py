#!/usr/bin/env python3
"""
Publication-Grade Thesis DOCX Generator (v2.0)
Extracts and builds native Word tables, centers figures with proper captions,
generates verified frontmatter with Level: 4, Semester: II, and stitches
everything into an official, flawless thesis.docx.
"""

import os
import sys
import re
import glob
import shutil
import subprocess
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls
from docxcompose.composer import Composer

def clean_cell_text(text):
    text = re.sub(r"\\textbf\{([^}]+)\}", r"\1", text)
    text = re.sub(r"\\mathbf\{([^}]+)\}", r"\1", text)
    text = re.sub(r"\\texttt\{([^}]+)\}", r"\1", text)
    text = re.sub(r"\\textit\{([^}]+)\}", r"\1", text)
    text = re.sub(r"\\text\{([^}]+)\}", r"\1", text)
    text = re.sub(r"\\mathbb\{R\}\^\{?(\d+)\}?", r"R^\1", text)
    text = re.sub(r"\\pm", r"±", text)
    text = re.sub(r"\\le", r"≤", text)
    text = re.sub(r"\\ge", r"≥", text)
    text = re.sub(r"\\times", r"×", text)
    text = re.sub(r"\\chi\^2", r"χ²", text)
    text = re.sub(r"\\Delta", r"Δ", text)
    text = re.sub(r"\\ll", r"≪", text)
    text = re.sub(r"\\mathcal\{([A-Za-z])\}", r"\1", text)
    text = re.sub(r"\\%", r"%", text)
    text = re.sub(r"\\_", r"_", text)
    text = re.sub(r"\\&", r"&", text)
    text = re.sub(r"---", r"—", text)
    text = re.sub(r"--", r"–", text)
    text = re.sub(r"\$([^$]+)\$", r"\1", text)
    text = text.replace("{", "").replace("}", "").replace("\\", "")
    return text.strip()

def set_cell_margins(cell, top=60, bottom=60, left=100, right=100):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def remove_table_borders(table):
    tblPr = table._tbl.tblPr
    for child in list(tblPr):
        if child.tag.endswith('tblBorders'):
            tblPr.remove(child)
    tblBorders = OxmlElement('w:tblBorders')
    for b in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        node = OxmlElement(f'w:{b}')
        node.set(qn('w:val'), 'none')
        tblBorders.append(node)
    tblPr.append(tblBorders)

def add_toc_field(paragraph):
    run = paragraph.add_run()
    fld1 = OxmlElement('w:fldChar')
    fld1.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld2 = OxmlElement('w:fldChar')
    fld2.set(qn('w:fldCharType'), 'separate')
    fld3 = OxmlElement('w:fldChar')
    fld3.set(qn('w:fldCharType'), 'end')
    run._r.append(fld1)
    run._r.append(instr)
    run._r.append(fld2)
    run._r.append(fld3)

def build_frontmatter(output_path):
    doc = docx.Document()
    
    # Configure A4 margins: Left 1.25", others 1.0"
    sec = doc.sections[0]
    sec.page_width = Inches(8.27)
    sec.page_height = Inches(11.69)
    sec.top_margin = Inches(1.0)
    sec.bottom_margin = Inches(1.0)
    sec.left_margin = Inches(1.25)
    sec.right_margin = Inches(1.0)

    def p(text="", size=11, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=4, line_spacing=1.15):
        par = doc.add_paragraph()
        par.alignment = align
        par.paragraph_format.space_before = Pt(space_before)
        par.paragraph_format.space_after = Pt(space_after)
        par.paragraph_format.line_spacing = line_spacing
        if text:
            r = par.add_run(text)
            r.font.name = "Times New Roman"
            r.font.size = Pt(size)
            r.bold = bold
            r.italic = italic
        return par

    # ─────────────────────────────────────────────────────────────
    # PAGE 1: COVER PAGE
    # ─────────────────────────────────────────────────────────────
    p("CertGraph: Heterogeneous Graph Attention Networks for Active Directory Certificate Services Vulnerability Detection and Autonomous Defense",
      size=17, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
      
    p("Course Code: ECE 452        Course Title: Project and Thesis",
      size=11.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
      
    p("Submitted By—", size=11.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=6)
    
    # Student IDs with Level 4, Semester II
    p("Student ID: 2002126        Level: 4, Semester: II", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    p("Student ID: 2002138        Level: 4, Semester: II", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    p("Student ID: 2102151        Level: 4, Semester: II", size=11, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    
    # HSTU Logo
    logo_par = doc.add_paragraph()
    logo_par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo_par.paragraph_format.space_before = Pt(4)
    logo_par.paragraph_format.space_after = Pt(10)
    if os.path.exists("figures/hstu_logo.png"):
        logo_par.add_run().add_picture("figures/hstu_logo.png", width=Inches(1.1))
        
    p("Submitted To—", size=11.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=4)
    p("Department of Electronics and Communication Engineering", size=12.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    p("in partial fulfillment of the requirements for the degree of", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    p("Bachelor of Science in Electronics and Communication Engineering", size=11.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=10)
    p("Hajee Mohammad Danesh Science and Technology University (HSTU)", size=12.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    p("Dinajpur-5200, Bangladesh", size=10.5, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=2)
    p("October, 2026", size=11.5, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    
    doc.add_page_break()

    # ─────────────────────────────────────────────────────────────
    # PAGE 2: CERTIFICATE
    # ─────────────────────────────────────────────────────────────
    p("Certificate", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)
    
    p("This is to certify that the thesis work entitled “CertGraph: Heterogeneous Graph Attention Networks for Active Directory Certificate Services Vulnerability Detection and Autonomous Defense” is carried by the following ID numbers: 2002126, 2002138, 2102151. To the fullest extent of our knowledge, we assert that this undertaking is an authentic and original contribution to the field. We certify that this thesis has not been previously submitted for the award of any other degree or diploma at this or any other institution.",
      size=11.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=16, line_spacing=1.2)
      
    p("Signed by the Final Examining committee:", size=11.5, bold=True, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=16)
    
    tbl_cert = doc.add_table(rows=7, cols=2)
    remove_table_borders(tbl_cert)
    tbl_cert.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_cert.autofit = False
    
    cert_data = [
        ("........................................", "........................................"),
        ("Chairman", "Supervisor"),
        ("Examination Committee", "Department of ECE, HSTU"),
        ("........................................", "........................................"),
        ("External Member", "Co-Supervisor"),
        ("Examination Committee", "Department of ECE, HSTU"),
        ("........................................\nInternal Member\nExamination Committee", "")
    ]
    
    for r_idx, (c1, c2) in enumerate(cert_data):
        row = tbl_cert.rows[r_idx]
        row.cells[0].width = Inches(2.9)
        row.cells[1].width = Inches(2.9)
        set_cell_margins(row.cells[0])
        set_cell_margins(row.cells[1])
        
        p1 = row.cells[0].paragraphs[0]
        p1.paragraph_format.line_spacing = 1.05
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(c1)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(10.5)
        if "Chairman" in c1 or "External" in c1 or "Internal" in c1:
            r1.bold = True
            
        p2 = row.cells[1].paragraphs[0]
        p2.paragraph_format.line_spacing = 1.05
        p2.paragraph_format.space_after = Pt(2)
        r2 = p2.add_run(c2)
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(10.5)
        if "Supervisor" in c2 or "Co-Supervisor" in c2:
            r2.bold = True

    doc.add_page_break()

    # ─────────────────────────────────────────────────────────────
    # PAGE 3: CANDIDATE'S DECLARATION
    # ─────────────────────────────────────────────────────────────
    p("Candidate's Declaration", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=12)
    
    p("We hereby declare that the research work presented in this thesis entitled “CertGraph: Heterogeneous Graph Attention Networks for Active Directory Certificate Services Vulnerability Detection and Autonomous Defense” is the outcome of an original investigation conducted by us under the supervision of the Department of Electronics and Communication Engineering, Hajee Mohammad Danesh Science and Technology University (HSTU), Dinajpur-5200, Bangladesh.",
      size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=8, line_spacing=1.15)
      
    p("We further solemnly declare that:", size=11, align=WD_ALIGN_PARAGRAPH.LEFT, space_after=4)
    
    p("1.  This work, or any part thereof, has not been submitted previously to any university or institution for the award of any degree, diploma, or other academic qualification.",
      size=10.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=2)
    p("2.  All material, concepts, and algorithms taken from the published or unpublished work of others have been fully and properly acknowledged and cited in accordance with standard academic referencing protocols.",
      size=10.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=2)
    p("3.  All synthetic datasets, forensic auditing scripts, empirical benchmark routines, and graph learning models described herein were constructed with rigorous adherence to ethical scientific standards and academic integrity.",
      size=10.5, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=12)
      
    tbl_decl = doc.add_table(rows=4, cols=2)
    remove_table_borders(tbl_decl)
    tbl_decl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_decl.autofit = False
    
    decl_data = [
        ("Date: October, 2026", "........................................\nStudent ID: 2002126  Level: 4, Semester: II"),
        ("Place: HSTU, Dinajpur", "........................................\nStudent ID: 2002138  Level: 4, Semester: II"),
        ("", "........................................\nStudent ID: 2102151  Level: 4, Semester: II"),
        ("", "Department of ECE, HSTU")
    ]
    
    for r_idx, (c1, c2) in enumerate(decl_data):
        row = tbl_decl.rows[r_idx]
        row.cells[0].width = Inches(2.9)
        row.cells[1].width = Inches(2.9)
        set_cell_margins(row.cells[0])
        set_cell_margins(row.cells[1])
        
        p1 = row.cells[0].paragraphs[0]
        p1.paragraph_format.line_spacing = 1.05
        p1.paragraph_format.space_after = Pt(2)
        r1 = p1.add_run(c1)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(10.5)
        if "Date" in c1 or "Place" in c1:
            r1.bold = True
            
        p2 = row.cells[1].paragraphs[0]
        p2.paragraph_format.line_spacing = 1.05
        p2.paragraph_format.space_after = Pt(2)
        r2 = p2.add_run(c2)
        r2.font.name = "Times New Roman"
        r2.font.size = Pt(10.5)

    doc.add_page_break()

    # ─────────────────────────────────────────────────────────────
    # PAGE 4: DEDICATION
    # ─────────────────────────────────────────────────────────────
    p("", space_after=140)
    p("Dedication", size=17, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=20)
    p("This thesis is dedicated to our beloved parents,\nwhose endless sacrifices, prayers, and unconditional love\nhave been the guiding light of our lives.",
      size=12, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=16, line_spacing=1.3)
    p("And to all our teachers and mentors,\nwho inspired our passion for computer science, cybersecurity, and scientific discovery.",
      size=12, italic=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0, line_spacing=1.3)
      
    doc.add_page_break()

    # ─────────────────────────────────────────────────────────────
    # PAGE 5: ACKNOWLEDGEMENTS
    # ─────────────────────────────────────────────────────────────
    p("Acknowledgements", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)
    
    p("We begin by thanking the Almighty for granting us the health, resolve, and clarity of thought to bring this undergraduate project and thesis to completion.",
      size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)
    p("We wish to express our heartfelt gratitude to our thesis supervisor and co-supervisor in the Department of Electronics and Communication Engineering at Hajee Mohammad Danesh Science and Technology University (HSTU). Throughout the research process—from initial problem formulation and mathematical proofs to the debugging of graph neural networks and manuscript preparation—their constructive criticism, technical feedback, and steady encouragement were vital to our progress.",
      size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)
    p("We also thank the Chairman and the faculty members of the Department of Electronics and Communication Engineering for their instruction, academic support, and the computing facilities provided to us over the course of our undergraduate studies.",
      size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)
    p("This work builds heavily on the contributions of the broader security and machine learning research communities. In particular, we acknowledge Will Schroeder and Lee Christensen (SpecterOps) for their foundational analysis of ADCS vulnerabilities; Oliver Lyak, Andy Robbins, and the BloodHound team for their offensive graph tools; the creators of the Game of Active Directory (GOAD) testing environment; and the maintainers of PyTorch Geometric and NetworkX.",
      size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6)
    p("Finally, we are deeply grateful to our parents and families. Their patience, moral support, and sacrifices made our education possible. We also thank our classmates, lab partners, and friends whose discussions, technical debates, and camaraderie helped us navigate the challenges of completing this degree.",
      size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=14)
      
    tbl_ack = doc.add_table(rows=2, cols=2)
    remove_table_borders(tbl_ack)
    tbl_ack.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl_ack.autofit = False
    
    ack_row0 = tbl_ack.rows[0]
    ack_row0.cells[0].width = Inches(2.9)
    ack_row0.cells[1].width = Inches(2.9)
    p_a1 = ack_row0.cells[0].paragraphs[0]
    p_a1.add_run("HSTU, Dinajpur\nOctober, 2026").bold = True
    p_a2 = ack_row0.cells[1].paragraphs[0]
    p_a2.add_run("Student ID: 2002126\nStudent ID: 2002138\nStudent ID: 2102151")
    
    ack_row1 = tbl_ack.rows[1]
    ack_row1.cells[0].width = Inches(2.9)
    ack_row1.cells[1].width = Inches(2.9)
    p_a4 = ack_row1.cells[1].paragraphs[0]
    p_a4.add_run("Department of ECE, HSTU")
    
    doc.add_page_break()

    # ─────────────────────────────────────────────────────────────
    # PAGE 6: ABSTRACT
    # ─────────────────────────────────────────────────────────────
    p("Abstract", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)
    
    p("Active Directory Certificate Services (ADCS) forms the identity and cryptographic backbone of modern Windows enterprise networks, issuing public-key certificates used for domain authentication, transport encryption, and single sign-on. When certificate templates, issuance policies, and access control lists (ACLs) are misconfigured, unprivileged accounts can obtain unauthorized administrative certificates and take over entire Active Directory forests. While signature-based scanners (such as Certipy and BloodHound) identify known vulnerable flag combinations, they either evaluate configuration attributes in isolation or incur high traversal costs on dense enterprise graphs.",
      size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, line_spacing=1.15)
      
    p("This thesis examines whether heterogeneous graph attention networks can reliably detect ADCS privilege escalation vectors (ESC1, ESC2, ESC3, ESC4, ESC9, and ESC13) directly from relational identity multigraphs. Using the experimental audit framework of Arp et al. (2022), we first evaluate synthetic Active Directory datasets and identify three systemic methodological pitfalls: positional identifier leakage in generation routines, feature asymmetry between baselines, and benchmark saturation where synthetic rules can be recovered by shallow tabular heuristics. We prove mathematically (Theorem 1) that residual skip-connections are indispensable for preventing representation collapse on identity graphs; without them, template representations become independent of their initial configuration flags, dropping classification F1 from 0.9986 to 0.4768.",
      size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, line_spacing=1.15)
      
    p("On an audited benchmark of 700 enterprise environments evaluated under 5-fold cross-validation with Nadeau-Bengio corrected variance estimators, our model achieves a Macro-F1 of 0.9986, significantly outperforming flat classifiers (0.8600) and signature heuristics (0.7791). However, when tested against an adversarial Zero-Shot Hard Negative suite (n = 63) containing non-exploitable templates with dangerous flags, pure neural models collapse to 1.59% accuracy due to shortcut learning on local template flags. Symbolic path traversal, by contrast, correctly recognizes path absence in 84.13% of cases.",
      size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=6, line_spacing=1.15)
      
    p("To reconcile this gap, we develop a Two-Tier Neuro-Symbolic architecture. The neural tier acts as a high-throughput candidate screener, filtering out over 89% of benign templates, while a localized symbolic oracle provides deterministic reachability proofs for flagged candidates. We show that false-positive suppression on hard negatives is achieved by design through this division of labor. Finally, we model identity privilege revocation as a Stackelberg security game and establish that optimal edge interdiction is NP-hard via reduction from Directed Multiway Cut, introducing a polynomial-time greedy heuristic that severs 92% of reachable attack paths under bounded operational budgets.",
      size=11, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=12, line_spacing=1.15)
      
    par_kw = p("", size=11, space_after=0)
    r_kwh = par_kw.add_run("Keywords: ")
    r_kwh.bold = True
    r_kwh.font.name = "Times New Roman"
    r_kwt = par_kw.add_run("Active Directory Certificate Services (ADCS), Graph Attention Networks, Heterogeneous Graphs, Identity and Access Management, Neuro-Symbolic Security.")
    r_kwt.font.name = "Times New Roman"

    doc.add_page_break()

    # ─────────────────────────────────────────────────────────────
    # PAGE 7: TABLE OF CONTENTS (Pre-rendered + Dynamic Field)
    # ─────────────────────────────────────────────────────────────
    p("Table of Contents", size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=14)
    
    toc_entries = [
        ("Chapter 1: Introduction", "1", True),
        ("    1.1 The Enterprise Identity Fabric and Perimeter Security", "1", False),
        ("    1.2 The Active Directory Certificate Services Vulnerability Landscape", "3", False),
        ("    1.3 Limitations of Current Auditing Methodologies", "5", False),
        ("    1.4 Research Objectives and Contributions", "7", False),
        ("    1.5 Thesis Organization", "8", False),
        ("Chapter 2: Background and the ADCS Threat Landscape", "9", True),
        ("    2.1 Architectural Taxonomy of Active Directory Certificate Services", "9", False),
        ("    2.2 Mathematical Model of Identity and PKI Graphs", "14", False),
        ("    2.3 Comprehensive ADCS Escalation Taxonomy (ESC1–ESC15+)", "18", False),
        ("Chapter 3: Formal Methodology and CertGraph Architecture", "26", True),
        ("    3.1 Mathematical Formalization of Enterprise Identity Multigraphs", "26", False),
        ("    3.2 The CertGraph Hetero-GAT Architecture", "30", False),
        ("    3.3 Loss Formulation and Training Dynamics", "38", False),
        ("Chapter 4: Forensic Audit of Security Machine Learning Pitfalls", "42", True),
        ("    4.1 Applying the Arp et al. Framework to Identity Graphs", "42", False),
        ("    4.2 Positional Index Leakage and Graph Canonicalization", "45", False),
        ("    4.3 Baseline Information Asymmetry and Feature Representation", "49", False),
        ("    4.4 Rule Saturation and the Limits of In-Distribution Benchmarks", "53", False),
        ("Chapter 5: Empirical Benchmarks and Comparative Evaluation", "58", True),
        ("    5.1 Experimental Setup and Dataset Curation", "58", False),
        ("    5.2 In-Distribution 5-Fold Cross-Validation", "62", False),
        ("    5.3 Systematic Architectural Ablation Studies", "66", False),
        ("    5.4 The Zero-Shot Adversarial Hard Negative Benchmark", "70", False),
        ("    5.5 Attention Weight Explainability and Graph Interpretability", "74", False),
        ("Chapter 6: Real-World Case Study: Game of Active Directory (GOAD)", "78", True),
        ("    6.1 GOAD Laboratory Multi-Forest Topology", "78", False),
        ("    6.2 Full-Chain Attack Emulation", "83", False),
        ("    6.3 CertGraph Deployment and End-to-End Auditing", "88", False),
        ("    6.4 Generalization to External Community Enterprise Topologies", "92", False),
        ("Chapter 7: Robustness, Scalability, and Deployment Considerations", "96", True),
        ("    7.1 Cross-Domain Generalization (ADSynth Benchmark)", "96", False),
        ("    7.2 Topological Noise Robustness and Incomplete Audits", "99", False),
        ("    7.3 Scalability, Memory Complexity, and Real-Time Inference", "102", False),
        ("Chapter 8: Neuro-Symbolic Hybridization for Identity Auditing", "106", True),
        ("    8.1 The Complementary Strengths of Neural and Symbolic Systems", "106", False),
        ("    8.2 The Two-Tier Neuro-Symbolic Pipeline", "108", False),
        ("    8.3 Hybrid Performance and Verification Overhead", "111", False),
        ("Chapter 9: Game-Theoretic Autonomous Defense and Edge Interdiction", "114", True),
        ("    9.1 The Stackelberg Security Game Formulation", "114", False),
        ("    9.2 Computational Complexity: NP-Hardness of Edge Interdiction", "117", False),
        ("    9.3 Greedy Approximation and Heuristic Interdiction Algorithms", "120", False),
        ("    9.4 Defensive Simulation and Budgeted Edge Blocking Evaluation", "123", False),
        ("Chapter 10: Conclusion and Future Directions", "126", True),
        ("    10.1 Summary of Contributions", "126", False),
        ("    10.2 Open Challenges and Limitations", "128", False),
        ("    10.3 Future Research Directions", "130", False),
        ("Mathematical Proof of Theorem 1", "132", True),
        ("Mathematical Proof of Theorem 2", "135", True),
        ("References", "138", True),
    ]

    for title, page_no, is_chap in toc_entries:
        par = doc.add_paragraph()
        par.paragraph_format.line_spacing = 1.05
        par.paragraph_format.space_before = Pt(3 if is_chap else 1)
        par.paragraph_format.space_after = Pt(2 if is_chap else 1)
        
        # Calculate dots leader
        dots_count = max(4, 75 - len(title) - len(page_no))
        dots = " ." * (dots_count // 2)
        
        r1 = par.add_run(title)
        r1.font.name = "Times New Roman"
        r1.font.size = Pt(10.5 if is_chap else 9.5)
        r1.bold = is_chap
        
        r_dots = par.add_run(dots)
        r_dots.font.name = "Times New Roman"
        r_dots.font.size = Pt(9)
        r_dots.font.color.rgb = RGBColor(160, 160, 160)
        
        r_p = par.add_run(f"  {page_no}")
        r_p.font.name = "Times New Roman"
        r_p.font.size = Pt(10 if is_chap else 9.5)
        r_p.bold = is_chap

    # Also add dynamic Word TOC field at bottom so Word users can press F9
    p("", space_after=10)
    toc_p = doc.add_paragraph()
    add_toc_field(toc_p)

    doc.add_page_break()
    doc.save(output_path)
    print(f"[✓] Created clean frontmatter: '{output_path}'")

def style_academic_table(table, col_widths, headers, rows):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Header row
    hdr_cells = table.rows[0].cells
    hdr_trPr = table.rows[0]._tr.get_or_add_trPr()
    hdr_trPr.append(parse_xml(r'<w:tblHeader {}/>'.format(nsdecls('w'))))
    hdr_trPr.append(parse_xml(r'<w:cantSplit {}/>'.format(nsdecls('w'))))
    
    for i, h in enumerate(headers):
        hdr_cells[i].width = col_widths[i]
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i > 0 else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.05
        r = p.add_run(h)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9.5)
        r.bold = True
        
        shd = parse_xml(r'<w:shd {} w:fill="F1F5F9"/>'.format(nsdecls('w')))
        hdr_cells[i]._tc.get_or_add_tcPr().append(shd)
        set_cell_margins(hdr_cells[i], top=60, bottom=60, left=80, right=80)

    # Data rows
    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx + 1]
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(r'<w:cantSplit {}/>'.format(nsdecls('w'))))
        row_cells = row.cells
        for c_idx, val in enumerate(row_data):
            if c_idx >= len(col_widths):
                break
            row_cells[c_idx].width = col_widths[c_idx]
            p = row_cells[c_idx].paragraphs[0]
            # Center if number/short text
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx > 0 and len(val) < 22 else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(val)
            r.font.name = 'Times New Roman'
            r.font.size = Pt(9.0)
            if "CertGraph" in val or "Reference" in val or "Full CertGraph" in val or "High Zero-Shot" in val or "100.0%" in val:
                if c_idx == 0:
                    r.bold = True
            set_cell_margins(row_cells[c_idx], top=50, bottom=50, left=80, right=80)

    # Border styling (classic booktabs style)
    tblPr = table._tbl.tblPr
    borders = parse_xml(
        r'<w:tblBorders {} >'
        r'  <w:top w:val="single" w:sz="8" w:space="0" w:color="334155"/>'
        r'  <w:bottom w:val="single" w:sz="8" w:space="0" w:color="334155"/>'
        r'  <w:insideH w:val="single" w:sz="4" w:space="0" w:color="E2E8F0"/>'
        r'  <w:left w:val="none"/>'
        r'  <w:right w:val="none"/>'
        r'  <w:insideV w:val="none"/>'
        r'</w:tblBorders>'.format(nsdecls('w'))
    )
    tblPr.append(borders)

def parse_latex_tables(tex_path, global_table_counter):
    """
    Extracts all table environments from a LaTeX file,
    returns modified text (with placeholders) and list of table dictionaries.
    """
    content = open(tex_path).read()
    tables = []
    
    def replacer(match):
        nonlocal global_table_counter
        raw = match.group(0)
        cap_m = re.search(r"\\caption\{([^}]+)\}", raw)
        caption = cap_m.group(1) if cap_m else "Table"
        
        # Extract between toprule and (bottomrule or end tabular)
        body_m = re.search(r"\\toprule(.*?)(?:\\bottomrule|\\end\{(?:tabular|tabularx)\})", raw, re.DOTALL)
        if not body_m:
            return ""
            
        tab_body = body_m.group(1)
        raw_lines = tab_body.split(r"\\")
        
        parsed_rows = []
        for l in raw_lines:
            l = re.sub(r"%.*$", "", l).strip()
            l = re.sub(r"\\(toprule|midrule|bottomrule|hline)", "", l).strip()
            if not l or "&" not in l:
                continue
            cells = [clean_cell_text(c) for c in l.split("&")]
            parsed_rows.append(cells)
            
        if not parsed_rows:
            return ""
            
        headers = parsed_rows[0]
        data_rows = parsed_rows[1:]
        
        # Calculate proportional column widths (total ~5.9 inches)
        col_count = len(headers)
        col_max_lens = [len(h) for h in headers]
        for row in data_rows:
            for c_idx, cell in enumerate(row):
                if c_idx < col_count:
                    col_max_lens[c_idx] = max(col_max_lens[c_idx], len(cell))
                    
        total_len = sum(col_max_lens)
        total_width = 5.9
        col_widths = []
        for l in col_max_lens:
            # Weighted width with min 0.7 inches
            w = max(0.7, (l / total_len) * total_width)
            col_widths.append(w)
            
        # Normalize to 5.9 inches
        current_sum = sum(col_widths)
        col_widths = [Inches((w / current_sum) * total_width) for w in col_widths]
        
        table_id = f"TABLE_PLACEHOLDER_{global_table_counter:03d}"
        global_table_counter += 1
        tables.append({
            "id": table_id,
            "caption": clean_cell_text(caption),
            "headers": headers,
            "rows": data_rows,
            "col_widths": col_widths
        })
        
        return f"\n\n\\par \\textbf{{[[[{table_id}]]]}} \\par\n\n"

    new_content = re.sub(r"\\begin\{table\}(?:\[[^\]]*\])?(.*?)\\end\{table\}", replacer, content, flags=re.DOTALL)
    return new_content, tables, global_table_counter

def main():
    repo_dir = "/home/hs32/Desktop/GOAD/thesis_paper"
    os.chdir(repo_dir)
    
    scratch_dir = "/tmp/thesis_build_v2"
    os.makedirs(scratch_dir, exist_ok=True)
    
    tex_inputs = [
        "chapters/ch01_introduction.tex",
        "chapters/ch02_background_threat.tex",
        "chapters/ch03_formal_methodology.tex",
        "chapters/ch04_forensic_audit.tex",
        "chapters/ch05_empirical_benchmarks.tex",
        "chapters/ch06_real_world_case_study.tex",
        "chapters/ch07_robustness_scalability.tex",
        "chapters/ch08_neuro_symbolic_hybrid.tex",
        "chapters/ch09_game_theoretic_defense.tex",
        "chapters/ch10_conclusion.tex",
        "proofs/theorem1_representation_collapse.tex",
        "proofs/theorem2_edge_blocking_nphardness.tex"
    ]
    
    print("[1/5] Preprocessing LaTeX chapters and extracting tables...")
    all_tables = {}
    preprocessed_files = []
    global_table_counter = 0
    
    for rel_path in tex_inputs:
        base_name = os.path.basename(rel_path)
        out_path = os.path.join(scratch_dir, base_name)
        new_text, tables, global_table_counter = parse_latex_tables(rel_path, global_table_counter)
        for t in tables:
            all_tables[t["id"]] = t
            
        # Promote proofs to chapters
        if "theorem1" in base_name:
            new_text = re.sub(r"\\section\[[^\]]*\]\{[^}]*\}", r"\\chapter{Appendix A: Mathematical Proof of Theorem 1 (Representation Preservation)}", new_text)
        elif "theorem2" in base_name:
            new_text = re.sub(r"\\section\[[^\]]*\]\{[^}]*\}", r"\\chapter{Appendix B: Mathematical Proof of Theorem 2 (NP-Hardness of Edge Interdiction)}", new_text)
            
        with open(out_path, "w") as f:
            f.write(new_text)
        preprocessed_files.append(out_path)
        
    print(f"      Extracted {len(all_tables)} tables for native Word rendering.")
    
    print("[2/5] Generating frontmatter via python-docx...")
    front_docx = os.path.join(scratch_dir, "frontmatter.docx")
    build_frontmatter(front_docx)
    
    print("[3/5] Compiling body chapters via Pandoc...")
    body_raw_docx = os.path.join(scratch_dir, "body_raw.docx")
    pandoc_cmd = [
        "/home/hs32/.local/bin/pandoc",
        *preprocessed_files,
        "--citeproc",
        "--bibliography=references.bib",
        "--reference-doc=old-reference-paper-to-follow-design/final(corrected).docx",
        "-o", body_raw_docx
    ]
    res = subprocess.run(pandoc_cmd, capture_output=True, text=True)
    if res.returncode != 0:
        print(f"[-] Pandoc failed:\n{res.stderr}")
        sys.exit(1)
        
    print("[4/5] Post-processing body: inserting native tables, centering figures & styling...")
    doc_body = docx.Document(body_raw_docx)
    
    # 1. Page breaks before Chapter Headings & Centering Figures
    for p in doc_body.paragraphs:
        if p.style.name.startswith("Heading 1"):
            p.paragraph_format.page_break_before = True
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(12)
        elif p.style.name.startswith("Heading 2"):
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            
        # Center graphics and keep with caption
        for r in p.runs:
            if 'graphic' in r._element.xml:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.keep_with_next = True
                
        # Style captions
        if p.style.name.lower().startswith('caption'):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(12)

    # 2. Insert References Heading 1 before the very first bibliography entry
    for p in doc_body.paragraphs:
        if "Arp, Daniel" in p.text:
            ref_heading = p.insert_paragraph_before()
            ref_heading.style = doc_body.styles["Heading 1"]
            ref_heading.text = "References"
            ref_heading.paragraph_format.page_break_before = True
            ref_heading.paragraph_format.space_before = Pt(18)
            ref_heading.paragraph_format.space_after = Pt(12)
            break

    # 3. Replace table placeholders with native Word tables
    p_idx = 0
    table_num = 1
    while p_idx < len(doc_body.paragraphs):
        p = doc_body.paragraphs[p_idx]
        match = re.search(r"\[\[\[(TABLE_PLACEHOLDER_\d+)\]\]\]", p.text)
        if match:
            t_id = match.group(1)
            t_data = all_tables.get(t_id)
            if t_data:
                # Insert caption above (keep with next so it stays on same page as table)
                cap_p = p.insert_paragraph_before()
                cap_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                cap_p.paragraph_format.space_before = Pt(14)
                cap_p.paragraph_format.space_after = Pt(6)
                cap_p.paragraph_format.keep_with_next = True
                
                r_c1 = cap_p.add_run(f"Table {table_num}: ")
                r_c1.bold = True
                r_c1.font.name = "Times New Roman"
                r_c1.font.size = Pt(10)
                r_c2 = cap_p.add_run(t_data["caption"])
                r_c2.font.name = "Times New Roman"
                r_c2.font.size = Pt(10)
                table_num += 1
                
                # Create table
                num_rows = len(t_data["rows"]) + 1
                num_cols = len(t_data["headers"])
                tbl = doc_body.add_table(rows=num_rows, cols=num_cols)
                style_academic_table(tbl, t_data["col_widths"], t_data["headers"], t_data["rows"])
                
                # Move table before placeholder paragraph
                p._p.addprevious(tbl._tbl)
                
                # Delete placeholder paragraph
                p._p.getparent().remove(p._p)
                p_idx -= 1
        p_idx += 1

    body_final_docx = os.path.join(scratch_dir, "body_final.docx")
    doc_body.save(body_final_docx)
    print(f"      Successfully styled body and inserted {table_num - 1} native Word tables.")

    print("[5/5] Stitching frontmatter and body via docxcompose...")
    master = docx.Document(front_docx)
    composer = Composer(master)
    composer.append(docx.Document(body_final_docx))
    
    final_output = "thesis.docx"
    composer.save(final_output)
    
    desktop_output = "/home/hs32/Desktop/thesis.docx"
    shutil.copyfile(final_output, desktop_output)
    
    size_mb = os.path.getsize(final_output) / (1024 * 1024)
    print(f"[✓] SUCCESS! Finished generating '{final_output}' ({size_mb:.2f} MB)")
    print(f"    Copied to Desktop at '{desktop_output}'")

if __name__ == "__main__":
    main()
