#!/usr/bin/env python3
"""
High-Fidelity DOCX Monograph Builder for CertGraph Thesis.
Combines Pandoc's native LaTeX/Math/Citation compiler with python-docx layout refinement
to match the exact official HSTU thesis format (replicated from final(corrected).docx).
"""

import os
import sys
import subprocess
import shutil
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def remove_table_borders(table):
    """Remove borders from a table for clean signature and metadata layout."""
    tblPr = table._tbl.tblPr
    for child in list(tblPr):
        if child.tag.endswith('tblBorders'):
            tblPr.remove(child)
            
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'none')
        tblBorders.append(border)
    tblPr.append(tblBorders)

def lock_table_widths(table, col_widths_in_inches):
    """Explicitly lock table and column widths to prevent cell wrapping."""
    total_dxa = int(sum(col_widths_in_inches) * 1440)
    
    # 1. Set tblW on tblPr
    tblPr = table._tbl.tblPr
    for child in list(tblPr):
        if child.tag.endswith('tblW'):
            tblPr.remove(child)
    tblW = OxmlElement('w:tblW')
    tblW.set(qn('w:w'), str(total_dxa))
    tblW.set(qn('w:type'), 'dxa')
    tblPr.append(tblW)
    
    # 2. Set tblGrid
    for child in list(table._tbl):
        if child.tag.endswith('tblGrid'):
            table._tbl.remove(child)
    tblGrid = OxmlElement('w:tblGrid')
    for w in col_widths_in_inches:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(int(w * 1440)))
        tblGrid.append(col)
    table._tbl.insert(1, tblGrid)
    
    # 3. Set cell widths on every cell
    for row in table.rows:
        trPr = row._tr.get_or_add_trPr()
        trPr.append(OxmlElement('w:cantSplit'))
        
        for c_idx, cell in enumerate(row.cells):
            if c_idx < len(col_widths_in_inches):
                w_dxa = str(int(col_widths_in_inches[c_idx] * 1440))
                tcPr = cell._tc.get_or_add_tcPr()
                for child in list(tcPr):
                    if child.tag.endswith('tcW'):
                        tcPr.remove(child)
                tcW = OxmlElement('w:tcW')
                tcW.set(qn('w:w'), w_dxa)
                tcW.set(qn('w:type'), 'dxa')
                tcPr.append(tcW)

def add_toc_field(paragraph):
    """Insert a native Microsoft Word TOC field into the paragraph."""
    run = paragraph.add_run()
    fldChar1 = OxmlElement('w:fldChar')
    fldChar1.set(qn('w:fldCharType'), 'begin')
    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = 'TOC \\o "1-3" \\h \\z \\u'
    fldChar2 = OxmlElement('w:fldChar')
    fldChar2.set(qn('w:fldCharType'), 'separate')
    fldChar3 = OxmlElement('w:fldChar')
    fldChar3.set(qn('w:fldCharType'), 'end')
    run._r.append(fldChar1)
    run._r.append(instrText)
    run._r.append(fldChar2)
    run._r.append(fldChar3)

def main():
    dir_path = os.path.dirname(os.path.abspath(__file__))
    os.chdir(dir_path)

    print("[1/3] Compiling chapters and references with Pandoc (citeproc + HSTU template)...")
    raw_docx = "/tmp/thesis_raw.docx"
    
    cmd = [
        "pandoc",
        "frontmatter_full.md",
        "chapters/ch01_introduction.tex",
        "chapters/ch02_background_threat.tex",
        "chapters/ch03_formal_methodology.tex",
        "chapters/ch04_forensic_audit.tex",
        "chapters/ch05_empirical_benchmarks.tex",
        "chapters/ch06_real_world_case_study.tex",
        "chapters/ch07_robustness_scalability.tex",
        "chapters/ch08_neuro_symbolic_hybrid.tex",
        "chapters/ch09_game_theoretic_defense.tex",
        "chapters/ch10_conclusion.tex",
        "proofs/theorem1_representation_collapse.tex",
        "proofs/theorem2_edge_blocking_nphardness.tex",
        "--citeproc",
        "--bibliography=references.bib",
        "--reference-doc=old-reference-paper-to-follow-design/final(corrected).docx",
        "-o", raw_docx
    ]
    
    res = subprocess.run(cmd, capture_output=True, text=True)
    if res.returncode != 0:
        print(f"[-] Pandoc failed:\n{res.stderr}")
        sys.exit(1)
        
    print("[2/3] Post-processing frontmatter, typography, and page breaks via python-docx...")
    doc = docx.Document(raw_docx)
    
    # 1. Standardize page margins (HSTU Thesis Specification: Left 1.25", others 1.0")
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.25)   # Left margin for binding
        section.right_margin = Inches(1.0)
        
    # 2. Configure signature tables (Certificate, Declaration, Acknowledgements)
    for i, table in enumerate(doc.tables[:3]):
        remove_table_borders(table)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        lock_table_widths(table, [2.85, 2.85])
        for row in table.rows:
            for cell in row.cells:
                tcPr = cell._tc.get_or_add_tcPr()
                tcMar = OxmlElement('w:tcMar')
                for m in ['top', 'bottom']:
                    node = OxmlElement(f'w:{m}')
                    node.set(qn('w:w'), '20')
                    node.set(qn('w:type'), 'dxa')
                    tcMar.append(node)
                for m in ['left', 'right']:
                    node = OxmlElement(f'w:{m}')
                    node.set(qn('w:w'), '60')
                    node.set(qn('w:type'), 'dxa')
                    tcMar.append(node)
                tcPr.append(tcMar)
                
                # Format text inside table cells
                for cp in cell.paragraphs:
                    cp.paragraph_format.line_spacing = 1.0
                    cp.paragraph_format.space_before = Pt(0)
                    cp.paragraph_format.space_after = Pt(1.5)
                    for cr in cp.runs:
                        cr.font.name = 'Times New Roman'
                        cr.font.size = Pt(10.5)
                
    # 3. Add explicit page breaks before all major sections / chapters
    frontmatter_breaks = [
        "certificate", "candidate's declaration", "dedication", 
        "acknowledgements", "abstract", "table of contents"
    ]
    
    current_section = "cover"
    for p in doc.paragraphs:
        txt = p.text.strip().lower()
        
        # Check section transitions
        for fb in frontmatter_breaks:
            if txt == fb:
                current_section = fb
                break
                
        # Format cover page
        if current_section == "cover":
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(2.0)
            p.paragraph_format.space_after = Pt(3.0)
            if p.text.strip():
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.name = 'Times New Roman'
                if txt.startswith("certgraph"):
                    for r in p.runs:
                        r.font.size = Pt(17)
                        r.bold = True
                        
        # Compact Certificate and Declaration paragraphs
        elif current_section in ["certificate", "candidate's declaration"]:
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(1.0)
            p.paragraph_format.space_after = Pt(2.5)
            for r in p.runs:
                r.font.name = 'Times New Roman'
                r.font.size = Pt(11)
                
        # Set page break before major headings
        if p.style.name.startswith("Heading 1"):
            if current_section != "cover":
                p.paragraph_format.page_break_before = True
        elif any(txt == fb for fb in frontmatter_breaks):
            p.paragraph_format.page_break_before = True
            
        # Center images and captions
        for r in p.runs:
            if 'graphic' in r._element.xml:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p.style.name.lower().startswith('caption'):
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            
        # Populate Table of Contents
        if txt == "table of contents":
            toc_p = doc.add_paragraph()
            add_toc_field(toc_p)
            p._p.addnext(toc_p._p)

    output_path = "thesis.docx"
    doc.save(output_path)
    
    # Also overwrite /home/hs32/Desktop/thesis.docx
    desktop_path = "/home/hs32/Desktop/thesis.docx"
    shutil.copyfile(output_path, desktop_path)
    
    size_mb = os.path.getsize(output_path) / (1024 * 1024)
    print(f"[3/3] Successfully generated '{output_path}' ({size_mb:.2f} MB)!")
    print(f"      Copied to '{desktop_path}'")

if __name__ == "__main__":
    main()
