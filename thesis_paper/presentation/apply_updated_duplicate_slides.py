#!/usr/bin/env python3
"""
apply_updated_duplicate_slides.py
Applies precise positioning and cleanup for:
- Slide 6 (05b/15): PROPOSED METHODOLOGY: GRAPH FORMULATION & PIPELINE
- Slide 8 (06b/15): BENCHMARK THREAT & PRIOR LIMITATIONS: ESC13
"""

import sys
import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

PPTX_PATH = "/home/hs32/Desktop/GOAD/thesis_paper/certgraph_thesis_defense.pptx"
OUT_PPTX = "/tmp/test_defense.pptx" if "--test" in sys.argv else PPTX_PATH

FIG_METHODOLOGY = "/home/hs32/Desktop/GOAD/thesis_paper/figures/methodology_gnn_pipeline.png"
FIG_ESC13 = "/home/hs32/Desktop/GOAD/thesis_paper/figures/esc13_benchmark_limitations.png"

prs = pptx.Presentation(PPTX_PATH)

# =============================================================================
# 1. UPDATE SLIDE 6 (05b/15)
# =============================================================================
slide6 = prs.slides[5]

shapes_to_remove6 = []
caption_shape6 = None

for s in slide6.shapes:
    if s.has_text_frame:
        txt = s.text_frame.text.strip()
        if "Figure:" in txt:
            caption_shape6 = s
            continue
        if any(k in txt for k in ["PROPOSED METHODOLOGY", "05b/15", "Course Title", "CertGraph GNN Pipeline", "Our methodology formalizes"]):
            continue
    if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
        if s.name in ['Group 2', 'Group 4', 'Group 9', 'Group 12']:
            continue
    shapes_to_remove6.append(s)

print(f"Slide 6: Removing {len(shapes_to_remove6)} old shapes...")
for s in shapes_to_remove6:
    sp = s._element
    sp.getparent().remove(sp)

# Add Diagram at top=3.82", height=4.70" (leaving plenty of room after subtitle)
pic6 = slide6.shapes.add_picture(FIG_METHODOLOGY, Inches(1.6), Inches(3.82), width=Inches(16.8), height=Inches(4.70))

if caption_shape6:
    caption_shape6.left = Inches(1.6)
    caption_shape6.top = Inches(8.65)
    caption_shape6.width = Inches(16.8)
    caption_shape6.text_frame.text = "Figure: End-to-End Proposed Methodology — Heterogeneous Multigraph Formalization, Inductive Hetero-GAT Engine, and Neuro-Symbolic Verification"
    p = caption_shape6.text_frame.paragraphs[0]
    p.font.name = "Times New Roman"
    p.font.size = Pt(17)
    p.font.italic = True
    p.alignment = PP_ALIGN.CENTER
    p.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

# =============================================================================
# 2. UPDATE SLIDE 8 (06b/15)
# =============================================================================
slide8 = prs.slides[7]

shapes_to_remove8 = []
caption_shape8 = None

for s in slide8.shapes:
    if s.has_text_frame:
        txt = s.text_frame.text.strip()
        if "Figure:" in txt:
            caption_shape8 = s
            continue
        if any(k in txt for k in ["BENCHMARK THREAT", "06b/15", "Course Title"]):
            continue
    # Keep only background (Group 2) and logo (Group 4)
    if s.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.GROUP:
        if s.name in ['Group 2', 'Group 4']:
            continue
    shapes_to_remove8.append(s)

print(f"Slide 8: Removing {len(shapes_to_remove8)} old shapes...")
for s in shapes_to_remove8:
    sp = s._element
    sp.getparent().remove(sp)

# Subtitle on Slide 8
sub8 = slide8.shapes.add_textbox(Inches(2.44), Inches(1.80), Inches(15.0), Inches(0.45))
tf8 = sub8.text_frame
tf8.word_wrap = True
tf8.text = "Evaluating state-of-the-art vulnerability auditing tools against composite multi-hop Active Directory attack paths."
p8 = tf8.paragraphs[0]
p8.font.name = "Times New Roman"
p8.font.size = Pt(18)
p8.font.italic = True
p8.font.color.rgb = RGBColor(0x1e, 0x29, 0x3b)

# Diagram at top=2.35", height=6.25"
pic8 = slide8.shapes.add_picture(FIG_ESC13, Inches(1.6), Inches(2.35), width=Inches(16.8), height=Inches(6.25))

if caption_shape8:
    caption_shape8.left = Inches(1.6)
    caption_shape8.top = Inches(8.75)
    caption_shape8.width = Inches(16.8)
    caption_shape8.text_frame.text = "Figure: Two-Hop ESC13 Attack Chain Evaluated as Our Core Benchmark — Tool Limitations vs. Proposed GNN Detection"
    p = caption_shape8.text_frame.paragraphs[0]
    p.font.name = "Times New Roman"
    p.font.size = Pt(17)
    p.font.italic = True
    p.alignment = PP_ALIGN.CENTER
    p.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

prs.save(OUT_PPTX)
print(f"[✓] Saved updated presentation to {OUT_PPTX}")
